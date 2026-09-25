"""Piloto local de siete imágenes; biblioteca persistente sin llamadas desde la cotización."""

from __future__ import annotations

import argparse
import html
import json
import os
import sqlite3
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

# Clientes y almacenamiento compartidos con el motor; el piloto conserva su CLI.
from mobiliti_saas.quote_engine.image_library import (
    CONFIGURACIONES, PROMPT, MAX_BYTES, huella, dimensiones, abrir_biblioteca,
    registrar_original, clave_version, reservar, completar, aprobar, obtener_aprobada,
    SinRedireccion, ErrorHTTPProveedor, solicitar, consultar_json, ejecutar_version,
    exportar_archivo,
)


def preparar(fuente: Path, db: sqlite3.Connection) -> dict:
    libro = load_workbook(fuente, data_only=True, keep_links=False)
    seleccion: dict[str, dict] = {}
    try:
        hoja = libro["Quotation"]
        for imagen in sorted(hoja._images, key=lambda i: i.anchor._from.row):
            datos = imagen._data()
            with Image.open(BytesIO(datos)) as original:
                if original.size != (150, 150) or original.format != "PNG":
                    continue
            sha = huella(datos)
            fila = imagen.anchor._from.row + 1
            if sha not in seleccion and len(seleccion) < 7:
                registrar_original(db, datos)
                seleccion[sha] = {"sha256": sha, "producto": str(hoja.cell(fila, 2).value or ""), "filas": []}
            if sha in seleccion:
                seleccion[sha]["filas"].append(fila)
    finally:
        libro.close()
    if len(seleccion) != 7:
        raise ValueError(f"El piloto necesita siete PNG únicos de 150 × 150; encontrados: {len(seleccion)}")
    return {"fuente": fuente.name, "sha256": huella(fuente.read_bytes()), "imagenes": list(seleccion.values()),
            "configuraciones": CONFIGURACIONES}


def generar_galeria(db: sqlite3.Connection, manifiesto: dict, salida: Path) -> Path:
    salida.mkdir(parents=True, exist_ok=True)
    paneles = []
    completadas = 0
    for numero, item in enumerate(manifiesto["imagenes"], 1):
        original = db.execute("SELECT imagen FROM originales WHERE sha=?", (item["sha256"],)).fetchone()[0]
        exportar_archivo(salida / (item["sha256"] + ".png"), original)
        celdas = [f'<figure><a href="{item["sha256"]}.png" target="_blank"><img src="{item["sha256"]}.png"></a><figcaption>Original · 150 × 150</figcaption></figure>']
        for proveedor, config in manifiesto["configuraciones"].items():
            identificador = clave_version(item["sha256"], proveedor, config)
            version = db.execute("SELECT * FROM versiones WHERE id=?", (identificador,)).fetchone()
            titulo = "GPT Image 2.5 Sunburst" if proveedor == "openai" else "SeedVR2"
            if version and version["estado"] == "completada":
                if huella(version["imagen"]) != version["sha"]:
                    raise ValueError("Resultado corrupto")
                exportar_archivo(salida / (identificador + ".png"), version["imagen"])
                detalles = json.loads(version["metadata"])
                exportar_archivo(salida / (identificador + ".json"), json.dumps({"id": identificador, "modelo": config["model"], **detalles}, indent=2).encode())
                ancho, alto = detalles["dimensiones"]
                celdas.append(f'<figure><a href="{identificador}.png" target="_blank"><img src="{identificador}.png"></a><figcaption>{titulo} · {ancho} × {alto} · <a href="{identificador}.json">Detalles</a></figcaption></figure>')
                completadas += 1
            else:
                estado = version["estado"] if version else "pendiente de generar"
                celdas.append(f'<figure><div class="pendiente">{titulo}<br>{html.escape(estado)}</div><figcaption>Sin resultado disponible</figcaption></figure>')
        paneles.append(f'<section><h2>{numero}. {html.escape(item["producto"])}</h2><p>Filas: {", ".join(map(str, item["filas"]))}</p><div class="fila">{"".join(celdas)}</div></section>')
    documento = '<!doctype html><html lang="es"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Comparativa GEELY</title><style>body{font:16px system-ui;margin:32px;background:#eef4f3;color:#173a37}main{max-width:1300px;margin:auto}section{background:white;padding:24px;margin:24px 0;border-radius:16px}.fila{display:grid;grid-template-columns:repeat(3,1fr);gap:20px}figure{margin:0}img,.pendiente{width:100%;aspect-ratio:1;object-fit:contain;background:white}.pendiente{display:grid;place-content:center;text-align:center;background:#f1f3f5;color:#596865}figcaption{padding:12px 0;font-size:14px}h2{font-size:20px}a{color:#00685c}@media(max-width:700px){.fila{grid-template-columns:1fr}}</style><main>'
    apariciones = sum(len(item["filas"]) for item in manifiesto["imagenes"])
    documento += f'<h1>Comparativa de imágenes GEELY</h1><p>{completadas}/14 resultados generados · 7 originales únicos · {apariciones} apariciones en Quotation.</p><p>Mismos originales en ambos proveedores. Haz clic en una imagen para verla a tamaño real. Revisa patas, ruedas, brazos, colores y geometría antes de aprobarla.</p>'
    documento += "".join(paneles) + '</main></html>'
    # La galería es una vista regenerable; los originales y resultados son inmutables.
    ruta = salida / "comparativa.html"
    ruta.write_text(documento, encoding="utf-8")
    return ruta


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fuente", type=Path, required=True)
    parser.add_argument("--biblioteca", type=Path, required=True, help="SQLite persistente, una biblioteca por cuenta")
    parser.add_argument("--salida", type=Path, required=True)
    parser.add_argument("--ejecutar", choices=["openai", "seedvr2", "ambos"], help="Sin esta opción no se llama a ninguna API")
    parser.add_argument("--aprobar", help="ID completo de la versión que ya fue revisada visualmente")
    args = parser.parse_args()
    db = abrir_biblioteca(args.biblioteca)
    errores = []
    try:
        manifiesto = preparar(args.fuente, db)
        args.salida.mkdir(parents=True, exist_ok=True)
        exportar_archivo(args.salida / "manifest.json", json.dumps(manifiesto, ensure_ascii=False, indent=2).encode())
        if args.aprobar:
            aprobar(db, args.aprobar)
        if args.ejecutar:
            proveedores = list(CONFIGURACIONES) if args.ejecutar == "ambos" else [args.ejecutar]
            for proveedor in proveedores:
                clave = os.environ.get("OPENAI_API_KEY" if proveedor == "openai" else "FAL_KEY", "").strip()
                if not clave:
                    errores.append(f"{proveedor}: falta la clave; no se hicieron solicitudes")
                    continue
                for numero, item in enumerate(manifiesto["imagenes"], 1):
                    print(f"{proveedor}: imagen {numero}/7", flush=True)
                    try:
                        ejecutar_version(db, item["sha256"], proveedor, CONFIGURACIONES[proveedor], clave)
                    except (ValueError, RuntimeError, OSError, KeyError) as error:
                        detalle = str(error) if isinstance(error, (ValueError, RuntimeError)) else type(error).__name__
                        errores.append(f"{proveedor}: {detalle}; detenido para revisar sin repetir cobros")
                        break
        print(generar_galeria(db, manifiesto, args.salida))
        for error in errores:
            print(error)
        return 2 if errores else 0
    finally:
        db.close()


if __name__ == "__main__":
    raise SystemExit(main())
